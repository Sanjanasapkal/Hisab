import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Theme Colors
    COLOR_PRIMARY = RGBColor(24, 59, 53)     # #183B35 (Deep Hisab Green)
    COLOR_ACCENT = RGBColor(31, 107, 87)     # #1F6B57 (Vibrant Forest Jade)
    COLOR_LIGHT_BG = RGBColor(246, 248, 247) # #F6F8F7
    COLOR_CARD = RGBColor(255, 255, 255)     # #FFFFFF
    COLOR_BORDER = RGBColor(220, 226, 224)   # Subtle Border
    COLOR_TEXT_DARK = RGBColor(28, 35, 33)   # Dark Charcoal
    COLOR_TEXT_MUTED = RGBColor(90, 105, 100)# Slate Muted
    COLOR_WHITE = RGBColor(255, 255, 255)
    COLOR_MINT = RGBColor(167, 211, 198)

    slides_data = [
        # Slide 1
        {
            "is_title": True,
            "title": "HISAB (हिसाब)",
            "subtitle": "Cloud-Synchronized Personal Finance & Ledger Management System",
            "presenter": "Developer: Sanjana Dadaso Sapkal",
            "meta": "Tech Stack: Android (Kotlin) | Node.js (Express) | MongoDB Atlas | Render Cloud"
        },
        # Slide 2
        {
            "num": "01",
            "title": "Problem Statement & Motivation",
            "tagline": "The limitations of traditional khata & informal lending tracking",
            "cards": [
                ("Fragile Traditional Bookkeeping", "Traditional paper-based ledger notebooks ('Bahi Khata') are vulnerable to physical damage, loss, ink fading, and human math calculation mistakes."),
                ("Disorganized Digital Alternatives", "Chat apps (WhatsApp/SMS) and note apps lack debit/credit logic, balance reconciliation, settlement audit trails, or multi-person balance summaries."),
                ("Overcomplicated Enterprise Tools", "Existing accounting tools (Tally, QuickBooks) are designed for complex business tax filings and are overly bulky for personal peer-to-peer lending."),
                ("The Hisab Vision", "A lightweight, secure, cloud-synced digital notebook specifically crafted for tracking loans, borrowed money, shared bills, and settlements effortlessly.")
            ]
        },
        # Slide 3
        {
            "num": "02",
            "title": "Project Objectives & Core Capabilities",
            "tagline": "Everything required for effortless peer debt management",
            "cards": [
                ("Dual-Party Financial Ledger", "Tracks both 'You will get' (+ Receivables) and 'You will give' (- Payables), calculating signed real-time net balances automatically."),
                ("Audit-Preserved Account Settlement", "'Settle Hisab' closes active debt cycles to ₹0 while preserving 100% of historical transactions for transparency and dispute prevention."),
                ("Enterprise Cloud Synchronization", "Continuous live sync between mobile devices and MongoDB Atlas cloud without manual exports or file backups."),
                ("In-App Security & Recovery", "Strict user data isolation, encrypted JWT sessions, and 6-digit email OTPs delivered for authentication and in-app password reset.")
            ]
        },
        # Slide 4
        {
            "num": "03",
            "title": "System Architecture & Data Flow",
            "tagline": "3-Tier cloud-connected enterprise mobile architecture",
            "cards": [
                ("Tier 1: Android Client (Kotlin)", "Built with Material Design 3, ViewBinding, Retrofit2, and Coroutines. Manages UI rendering, offline-ready session caching, and secure token headers."),
                ("Tier 2: Cloud Backend (Node.js/Express)", "Hosted 24/7 on Render cloud. Enforces RESTful standards, route rate limiting, input sanitization, and JWT authentication middleware."),
                ("Tier 3: Database (MongoDB Atlas)", "Cloud-hosted NoSQL cluster with Mongoose ORM. Enforces relational schema validation, index optimization, and strict multi-tenant user scoping."),
                ("Communication & Notifications", "TLS-encrypted HTTPS over public internet. Brevo SMTP Relay via cloud-safe port 2525 for instantaneous 6-digit email OTP delivery.")
            ]
        },
        # Slide 5
        {
            "num": "04",
            "title": "Android Application Architecture & UI/UX",
            "tagline": "Modern, intuitive mobile interface designed for daily usability",
            "cards": [
                ("Visual Design & Branding", "Custom deep-green color scheme (#183B35), card elevations, vector iconography, dedicated splash screen, and dynamic rupee formatting (₹)."),
                ("Smart Dashboard (MainActivity)", "Interactive 3-part financial summary card (Receivables, Payables, Net Balance) with real-time person search and instant add modal."),
                ("Account Ledger (AccountDetailActivity)", "Live signed status badges (Owes You / You Owe / Settled), reverse-chronological transaction stream, and one-tap 'Settle Hisab' dialog."),
                ("Audit Trail (TransactionHistoryActivity)", "Historical audit screen displaying past settlement cycles with payment modes (UPI/Cash), settlement notes, and timestamps.")
            ]
        },
        # Slide 6
        {
            "num": "05",
            "title": "Cloud Backend & API Architecture",
            "tagline": "Scalable REST API endpoints engineered for security and reliability",
            "cards": [
                ("Auth & In-App Recovery APIs", "POST /api/auth/register, verify-email-otp, login with JWT, and in-app forgot/reset-password with 6-digit codes and zero web URLs."),
                ("People & Contact Management", "Full CRUD operations on /api/people with automatic balance computation, duplicate detection, and cascade deletion guards."),
                ("Transactions & Settlement Engine", "POST /api/transactions to record debts or credits. POST /api/transactions/settle to archive cycles into AccountPeriods atomically."),
                ("OTA Version Updates & Release News", "GET /api/app/version delivers live version checks, 'What's New' highlights, and one-tap APK updates straight from MongoDB.")
            ]
        },
        # Slide 7
        {
            "num": "06",
            "title": "Security & Data Privacy Implementation",
            "tagline": "Defensive security safeguards protecting sensitive financial records",
            "cards": [
                ("Multi-Tenant Data Isolation", "Every database query enforces { userId: req.user._id }. User A can never inspect, alter, or delete User B's financial records."),
                ("Cryptographic Protection", "Bcrypt password hashing (salt rounds: 10), SHA-256 OTP hashing in database, and cryptographic random token generation."),
                ("Anti-Enumeration & Rate Limiting", "Authentication endpoints return generic success responses to prevent username/email harvesting; IP rate-limiters prevent brute force."),
                ("Cloud Firewall Hardening", "Custom port 2525 routing to bypass cloud host port 587 outbound firewall restrictions on Render's free tier.")
            ]
        },
        # Slide 8
        {
            "num": "07",
            "title": "Testing, Validation & Quality Assurance",
            "tagline": "Rigorous automated testing guaranteeing zero defects in production",
            "cards": [
                ("100% Passing Automated Suite", "73 out of 73 unit and integration tests passing with 0 failures using Node.js native test runner and Supertest."),
                ("Strict User Data Isolation Tests", "Automated attack simulation tests verifying that unauthorized cross-user queries return 404 Not Found as expected."),
                ("Financial Logic Verification", "Mathematical precision tests proving balance addition, negative subtraction, signed balances, and settlement zeroing."),
                ("Provider & Resiliency Testing", "Mock email fallback testing, SMTP error recovery, cloud database reconnection, and network retry validation.")
            ]
        },
        # Slide 9
        {
            "num": "08",
            "title": "Cloud Deployment & DevOps Pipeline",
            "tagline": "Zero local dependency — accessible anywhere, anytime, globally",
            "cards": [
                ("Continuous Deployment (CI/CD)", "GitHub-integrated automated pipeline. Every git push to main automatically triggers a zero-downtime build on Render cloud."),
                ("Cloud Database Infrastructure", "MongoDB Atlas 3-node replica set with automated failover, TLS encryption in-transit, and AES-256 encryption at-rest."),
                ("Global Anywhere Access", "Works on any Wi-Fi, hotspot, or mobile cellular data (4G/5G). Does not require the developer's laptop to stay online."),
                ("Brevo Cloud Mail Gateway", "Configured cloud relay with SPF/DKIM verification ensuring high inbox delivery rate for OTP emails.")
            ]
        },
        # Slide 10
        {
            "num": "09",
            "title": "Conclusion & Future Scope",
            "tagline": "Summary of deliverables and strategic technical roadmap",
            "cards": [
                ("Key Deliverables Achieved", "A robust, fully functional, cloud-synced financial ledger app with production cloud deployment, secure email delivery, and clean modern UX."),
                ("Instant UPI Deep Linking", "Future integration with UPI payment apps (Google Pay, PhonePe, Paytm) to settle balances directly within the application."),
                ("PDF & Excel Report Exports", "Ability to export complete monthly or annual transaction summaries into stamped PDF and CSV statements."),
                ("Group Expense Splitting", "Extending beyond 1-on-1 ledgers to group trip and roommate expense splitting with automatic debt simplification.")
            ]
        }
    ]

    for item in slides_data:
        slide = prs.slides.add_slide(blank_layout)

        if item.get("is_title"):
            # Background Header block
            bg_shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
            bg_shape.fill.solid()
            bg_shape.fill.fore_color.rgb = COLOR_PRIMARY
            bg_shape.line.fill.background()

            # Decorative inner box
            card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.2), Inches(1.0), Inches(10.933), Inches(5.5))
            card.fill.solid()
            card.fill.fore_color.rgb = RGBColor(20, 50, 45)
            card.line.color.rgb = COLOR_ACCENT
            card.line.width = Pt(2)

            # Title
            txBox = slide.shapes.add_textbox(Inches(1.8), Inches(1.6), Inches(9.733), Inches(1.5))
            tf = txBox.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = item["title"]
            p.font.size = Pt(48)
            p.font.bold = True
            p.font.color.rgb = COLOR_WHITE
            p.alignment = PP_ALIGN.CENTER

            # Subtitle
            txBox_sub = slide.shapes.add_textbox(Inches(1.8), Inches(3.0), Inches(9.733), Inches(1.0))
            tf_sub = txBox_sub.text_frame
            tf_sub.word_wrap = True
            p_sub = tf_sub.paragraphs[0]
            p_sub.text = item["subtitle"]
            p_sub.font.size = Pt(22)
            p_sub.font.color.rgb = COLOR_MINT
            p_sub.alignment = PP_ALIGN.CENTER

            # Divider line
            div = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(4.5), Inches(4.2), Inches(4.333), Inches(0.04))
            div.fill.solid()
            div.fill.fore_color.rgb = COLOR_ACCENT
            div.line.fill.background()

            # Presenter
            txBox_pres = slide.shapes.add_textbox(Inches(1.8), Inches(4.5), Inches(9.733), Inches(0.8))
            tf_pres = txBox_pres.text_frame
            p_pres = tf_pres.paragraphs[0]
            p_pres.text = item["presenter"]
            p_pres.font.size = Pt(18)
            p_pres.font.bold = True
            p_pres.font.color.rgb = COLOR_WHITE
            p_pres.alignment = PP_ALIGN.CENTER

            # Meta / Tech Stack
            txBox_meta = slide.shapes.add_textbox(Inches(1.8), Inches(5.2), Inches(9.733), Inches(0.8))
            tf_meta = txBox_meta.text_frame
            p_meta = tf_meta.paragraphs[0]
            p_meta.text = item["meta"]
            p_meta.font.size = Pt(14)
            p_meta.font.color.rgb = RGBColor(180, 200, 195)
            p_meta.alignment = PP_ALIGN.CENTER

        else:
            # Light slide background
            bg_shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
            bg_shape.fill.solid()
            bg_shape.fill.fore_color.rgb = COLOR_LIGHT_BG
            bg_shape.line.fill.background()

            # Top Header Bar
            top_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(1.35))
            top_bar.fill.solid()
            top_bar.fill.fore_color.rgb = COLOR_PRIMARY
            top_bar.line.fill.background()

            # Number badge
            badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.35), Inches(0.85), Inches(0.65))
            badge.fill.solid()
            badge.fill.fore_color.rgb = COLOR_ACCENT
            badge.line.fill.background()
            b_tf = badge.text_frame
            b_p = b_tf.paragraphs[0]
            b_p.text = item["num"]
            b_p.font.size = Pt(18)
            b_p.font.bold = True
            b_p.font.color.rgb = COLOR_WHITE
            b_p.alignment = PP_ALIGN.CENTER

            # Slide Title in Header
            tx_title = slide.shapes.add_textbox(Inches(1.85), Inches(0.2), Inches(10.5), Inches(0.6))
            tf_t = tx_title.text_frame
            p_t = tf_t.paragraphs[0]
            p_t.text = item["title"]
            p_t.font.size = Pt(24)
            p_t.font.bold = True
            p_t.font.color.rgb = COLOR_WHITE

            # Tagline in Header
            tx_tag = slide.shapes.add_textbox(Inches(1.85), Inches(0.75), Inches(10.5), Inches(0.45))
            tf_tag = tx_tag.text_frame
            p_tag = tf_tag.paragraphs[0]
            p_tag.text = item["tagline"]
            p_tag.font.size = Pt(13)
            p_tag.font.color.rgb = COLOR_MINT

            # Render 4 Grid Cards (2x2 layout)
            cards = item.get("cards", [])
            card_w = Inches(5.6)
            card_h = Inches(2.45)
            positions = [
                (Inches(0.8), Inches(1.75)),   # Card 1 (Top Left)
                (Inches(6.85), Inches(1.75)),  # Card 2 (Top Right)
                (Inches(0.8), Inches(4.55)),   # Card 3 (Bottom Left)
                (Inches(6.85), Inches(4.55))   # Card 4 (Bottom Right)
            ]

            for idx, (c_title, c_desc) in enumerate(cards):
                left, top = positions[idx]
                card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, card_w, card_h)
                card.fill.solid()
                card.fill.fore_color.rgb = COLOR_CARD
                card.line.color.rgb = COLOR_BORDER
                card.line.width = Pt(1.5)

                # Accent line on left of card
                accent_strip = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top + Inches(0.2), Inches(0.08), Inches(2.05))
                accent_strip.fill.solid()
                accent_strip.fill.fore_color.rgb = COLOR_ACCENT
                accent_strip.line.fill.background()

                # Text box inside card
                tb = slide.shapes.add_textbox(left + Inches(0.25), top + Inches(0.15), card_w - Inches(0.4), card_h - Inches(0.3))
                ctf = tb.text_frame
                ctf.word_wrap = True

                # Card Heading
                cp1 = ctf.paragraphs[0]
                cp1.text = c_title
                cp1.font.size = Pt(16)
                cp1.font.bold = True
                cp1.font.color.rgb = COLOR_PRIMARY
                cp1.space_after = Pt(8)

                # Card Description
                cp2 = ctf.add_paragraph()
                cp2.text = c_desc
                cp2.font.size = Pt(12.5)
                cp2.font.color.rgb = COLOR_TEXT_DARK
                cp2.line_spacing = 1.25

            # Bottom Footer
            foot_tb = slide.shapes.add_textbox(Inches(0.8), Inches(7.05), Inches(11.65), Inches(0.35))
            ftf = foot_tb.text_frame
            fp = ftf.paragraphs[0]
            fp.text = "Hisab Project Presentation • Sanjana Dadaso Sapkal • Deployed on Render & MongoDB Atlas"
            fp.font.size = Pt(10)
            fp.font.color.rgb = COLOR_TEXT_MUTED

    output_path = r"c:\Users\HP\AndroidStudioProjects\Hisab\Hisab_Project_Presentation.pptx"
    prs.save(output_path)
    print(f"SUCCESS: Presentation saved at: {output_path}")

if __name__ == "__main__":
    create_presentation()
